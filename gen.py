import random, math, datetime
from openpyxl import Workbook
from openpyxl.styles import Font
from docx import Document
from docx.shared import Pt

civ="Internal Revenue Service,Social Security Administration,Department of Transportation,Department of Justice,Department of the Interior,Department of Homeland Security,United States Postal Service".split(",")
dfn="Army,Navy,Marines,Air Force,Space Force".split(",")
states="Alabama Alaska Arizona Arkansas California Colorado Connecticut Delaware Florida Georgia Hawaii Idaho Illinois Indiana Iowa Kansas Kentucky Louisiana Maine Maryland Massachusetts Michigan Minnesota Mississippi Missouri Montana Nebraska Nevada|New_Hampshire|New_Jersey|New_Mexico|New_York|North_Carolina|North_Dakota Ohio Oklahoma Oregon Pennsylvania|Rhode_Island|South_Carolina|South_Dakota Tennessee Texas Utah Vermont Virginia Washington|West_Virginia Wisconsin Wyoming".replace("|"," ").split()
states=["State of "+s.replace("_"," ") for s in states]
assert len(states)==50
cities="New York, NY;Los Angeles, CA;Chicago, IL;Houston, TX;Phoenix, AZ;Philadelphia, PA;San Antonio, TX;San Diego, CA;Dallas, TX;Jacksonville, FL;Austin, TX;Fort Worth, TX;San Jose, CA;Columbus, OH;Charlotte, NC;Indianapolis, IN;San Francisco, CA;Seattle, WA;Denver, CO;Oklahoma City, OK;Nashville, TN;El Paso, TX;Washington, DC;Boston, MA;Las Vegas, NV;Detroit, MI;Portland, OR;Memphis, TN;Louisville, KY;Milwaukee, WI;Baltimore, MD;Albuquerque, NM;Tucson, AZ;Fresno, CA;Sacramento, CA;Mesa, AZ;Atlanta, GA;Kansas City, MO;Colorado Springs, CO;Omaha, NE;Raleigh, NC;Virginia Beach, VA;Long Beach, CA;Miami, FL;Oakland, CA;Minneapolis, MN;Tulsa, OK;Bakersfield, CA;Tampa, FL;Arlington, TX".split(";")
ents=[(e,"Federal Civilian") for e in civ]+[(e,"Federal Defense") for e in dfn]+[(e,"State and Local") for e in states+cities]
first="James Mary Robert Patricia John Jennifer Michael Linda David Elizabeth William Barbara Richard Susan Joseph Jessica Thomas Sarah Christopher Karen Daniel Nancy Matthew Lisa Anthony Betty Mark Margaret Steven Sandra Andrew Ashley Kevin Emily Brian Donna Jason Michelle".split()
last="Smith Johnson Williams Brown Jones Garcia Miller Davis Rodriguez Martinez Hernandez Lopez Gonzalez Wilson Anderson Thomas Taylor Moore Jackson Martin Lee Perez Thompson White Harris Sanchez Clark Ramirez Lewis Robinson Walker Young Allen King Wright Scott Torres Nguyen Hill Flores".split()
types=["Greenfield security customer","Greenfield Observability customer","On prem to cloud migration","Customer Renewal","Customer expansion"]
stages=["Prospecting","Discovery","Proposal","Evaluation","Negotiation","Procurement","Closed Won","Closed Lost","Delayed"]
prods="SIEM SOAR Observe DEM WAF PlatformOnPrem PlatformCloud Work".split()
acts=["first call","to be scheduled","Proof of Concept","Discovery Call","Requirements definition","solution creation","demonstration","bakeoff","engage Product dev"]
gaps=["feature missing","poor performance","none"]
comps="Minisoft MobStrike Dynatrack Goggles Amaze OldRelic DataBlocks InElastic IBN HPF Other".split()
# type -> plausible products
tp={types[0]:["SIEM","SOAR","WAF"],types[1]:["Observe","DEM"],types[2]:["PlatformCloud","PlatformOnPrem","Work"],types[3]:prods,types[4]:prods}
chosen=random.sample(ents,100) if len(ents)>=100 else None
rows=[]
for e,m in chosen:
    t=random.choice(types); st=random.choice(stages)
    early=stages.index(st)<3 if st in stages[:6] else False
    price=min(100_000_000,int(round(math.exp(random.gauss(13.3,1.3)),-3)))
    price=max(price,25000)
    win=random.choices(["Yes","No"],[38,62])[0]
    if st=="Closed Won": win="Yes"
    if early: win="No" if random.random()<.85 else "Yes"
    if win=="No" and random.random()<.55: gap=random.choice(gaps[:2])
    else: gap=random.choices(gaps,[10,10,80])[0]
    rows.append([e,m,f"{random.choice(first)} {random.choice(last)}",price,random.choice([0,0,5,10,10,15,20,25,30,40]),t,st,random.choice(tp[t]),random.choice(acts),random.choice(acts),gap,random.choice(comps),win])
hdr=["Government Entity Name","Public Sector Market","Customer Point of Contact","Total List Price of Amazing Software Solutions (USD)","Discount (%)","Opportunity Type","Deal Stage","Amazing Software Products Proposed","Solutions Engineer Recent Activity","Solutions Engineer Next Activity","Product Gap","Competitor","Technical Win"]
wb=Workbook(); ws=wb.active; ws.title="Opportunities"; ws.append(hdr)
for c in ws[1]: c.font=Font(bold=True)
for r in rows: ws.append(r)
for i,w in enumerate([30,20,24,26,10,32,14,22,24,24,18,14,13]): ws.column_dimensions[chr(65+i)].width=w
for r in range(2,102): ws.cell(r,4).number_format='$#,##0'
ws.freeze_panes="A2"; wb.save("SFDC_Opportunities_Export.xlsx")
no=[r for r in rows if r[12]=="No"]; tot=sum(r[3] for r in no)
top=sorted(no,key=lambda r:-r[3])[:15]
d=Document(); d.add_heading("Summary of Public Sector Sales Opportunities",0)
t=datetime.date.today(); d.add_paragraph(f"Date - {t:%B} {t.day}, {t.year}")
d.add_paragraph(f"Total number of sales opportunities - {len(rows)}\nTotal number with Technical Win = No - {len(no)}\nTotal revenue of deals without a Technical Win - ${tot:,}")
d.add_heading("Top 15 Deals without a Technical Win",1)
for r in top:
    parts=[r[0],r[1],f"${r[3]:,}",r[5],r[9]]
    if r[10]!="none": parts.append(r[10])
    parts.append(r[11]); d.add_paragraph(", ".join(parts))
d.save("SFDC_Summary_Report.docx")
print(len(rows),len(no),tot); print("\n".join(", ".join(map(str,[r[0],r[1],r[3],r[10]])) for r in top[:3]))
